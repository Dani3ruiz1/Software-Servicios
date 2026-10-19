"""Fase 5 contra Postgres real: ventas, caja, cuadre, concurrencia y reportes cuadrando con el libro."""
import asyncio
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from uuid import UUID
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from test_phase4_flow import _customer, _setup, _staff_with_schedule, _tomorrow_10

from app.core.db import scoped_session
from app.core.errors import AppError
from app.core.tenancy import Principal
from app.modules.appointments import service as appts
from app.modules.appointments.schemas import AppointmentIn, StatusIn
from app.modules.cash import service as cash
from app.modules.cash.schemas import CloseIn, MovementIn, OpenIn
from app.modules.catalog import service as catalog
from app.modules.catalog.schemas import ServiceIn
from app.modules.reports import export
from app.modules.reports import service as reports
from app.modules.sales import service as sales
from app.modules.sales.schemas import PayIn, PaymentIn, SaleIn, SaleItemIn, SendIn, VoidIn

D = Decimal
EF = lambda n: PaymentIn(metodo="efectivo", monto=D(n))  # noqa: E731
TZ = ZoneInfo("America/Bogota")


async def _svc(p: Principal, nombre: str, precio: int, *, editable=False, iva=False) -> UUID:
    async with scoped_session(tenant_id=p.tenant_id) as db:
        return (await catalog.create_service(db, p, ServiceIn(nombre=nombre, duracion_min=30, precio=D(precio),
                                                              precio_editable=editable, aplica_iva=iva))).id


async def _open_cash(p: Principal, inicial=0):
    async with scoped_session(tenant_id=p.tenant_id) as db:
        return await cash.open_session(db, p, OpenIn(monto_inicial=D(inicial)))


async def _sell(p: Principal, body: SaleIn):
    async with scoped_session(tenant_id=p.tenant_id) as db:
        return await sales.create_sale(db, p, body)


async def _appt(p: Principal, cid: UUID, svc: UUID, when: datetime | None = None, staff: UUID | None = None):
    async with scoped_session(tenant_id=p.tenant_id) as db:
        return await appts.create_appointment(db, p, AppointmentIn(customer_id=cid, inicio=when or _tomorrow_10(),
                                                                   service_ids=[svc], staff_id=staff))


def _today():
    return datetime.now(TZ).date()


async def test_cash_required_to_collect_but_not_to_open_a_tab():
    p = await _setup("Básico")
    svc = await _svc(p, "Corte", 30000)
    tab = await _sell(p, SaleIn(items=[SaleItemIn(service_id=svc)]))                  # abierta: no exige caja
    assert tab.estado == "abierta" and tab.comprobante_numero is None
    with pytest.raises(AppError) as e:
        await _sell(p, SaleIn(items=[SaleItemIn(service_id=svc)], pagos=[EF(30000)]))
    assert e.value.code == "CASH_SESSION_REQUIRED"
    await _open_cash(p)
    async with scoped_session(tenant_id=p.tenant_id) as db:
        paid = await sales.pay_sale(db, p, tab.id, PayIn(pagos=[EF(30000)]))
    assert paid.estado == "pagada" and paid.comprobante_numero == "CI-000001" and paid.cash_session_id


async def test_independent_plan_sells_without_cash_session():
    p = await _setup("Independiente")                                                  # sin módulo de caja
    svc = await _svc(p, "Corte", 20000)
    s = await _sell(p, SaleIn(items=[SaleItemIn(service_id=svc)], pagos=[EF(20000)]))
    assert s.estado == "pagada" and s.cash_session_id is None


async def test_full_flow_split_payment_cuadre_close_and_lock_after_close():
    p = await _setup("Básico")
    svc = await _svc(p, "Corte", 30000)
    cid = await _customer(p)
    ses = await _open_cash(p, 50000)
    with pytest.raises(AppError) as e:
        await _open_cash(p)
    assert e.value.code == "CASH_SESSION_ALREADY_OPEN"

    appt = await _appt(p, cid, svc)
    s1 = await _sell(p, SaleIn(appointment_id=appt.id, propina=D(2000), pagos=[
        EF(20000), PaymentIn(metodo="transferencia", monto=D(12000), entidad_bancaria="Bancolombia", referencia="TX1")]))
    assert s1.total == D("32000") and s1.comprobante_numero == "CI-000001" and s1.cash_session_id == ses.id
    assert s1.customer_id == cid and {x.metodo for x in s1.pagos} == {"efectivo", "transferencia"}
    async with scoped_session(tenant_id=p.tenant_id) as db:
        assert (await appts.get_appointment(db, p, appt.id)).estado == "atendida"      # cobrar = atendida
        with pytest.raises(AppError) as e:                                              # no se cobra dos veces la cita
            await sales.create_sale(db, p, SaleIn(appointment_id=appt.id, pagos=[EF(30000)]))
        assert e.value.code == "CONFLICT"

    s2 = await _sell(p, SaleIn(items=[SaleItemIn(descripcion="Gaseosa", precio_cobrado=D(5000))], pagos=[EF(5000)]))
    assert s2.comprobante_numero == "CI-000002"                                         # consecutivo sin huecos

    async with scoped_session(tenant_id=p.tenant_id) as db:
        await cash.add_movement(db, p, MovementIn(tipo="egreso", categoria="insumos", monto=D(3000)))
        c = await cash.cuadre(db, p, ses.id)
    assert c.ventas_efectivo == D("25000") and c.egresos == D("3000") and c.esperado_efectivo == D("72000")
    assert {m.metodo: m.monto for m in c.por_metodo} == {"efectivo": D("25000"), "transferencia": D("12000")}
    assert c.ventas_pagadas == 2 and c.propinas == D("2000") and c.total_vendido == D("37000")

    async with scoped_session(tenant_id=p.tenant_id) as db:
        closed = await cash.close_session(db, p, ses.id, CloseIn(monto_contado=D(71000)))
    assert closed.esperado_efectivo == D("72000") and closed.diferencia == D("-1000") and closed.sesion.estado == "cerrada"

    with pytest.raises(AppError) as e:
        async with scoped_session(tenant_id=p.tenant_id) as db:
            await sales.void_sale(db, p, s2.id, VoidIn(motivo="error de digitación"))
    assert e.value.code == "CONFLICT"                                                   # su caja ya cerró
    with pytest.raises(AppError) as e:
        await _sell(p, SaleIn(items=[SaleItemIn(descripcion="X", precio_cobrado=D(1000))], pagos=[EF(1000)]))
    assert e.value.code == "CASH_SESSION_REQUIRED"
    with pytest.raises(AppError):
        async with scoped_session(tenant_id=p.tenant_id) as db:
            await cash.close_session(db, p, ses.id, CloseIn(monto_contado=D(1)))        # no se cierra dos veces


async def test_price_editable_and_original_price_is_always_kept():
    p = await _setup("Básico")
    fijo, libre = await _svc(p, "Fijo", 30000), await _svc(p, "Libre", 40000, editable=True)
    await _open_cash(p)
    with pytest.raises(AppError) as e:
        await _sell(p, SaleIn(items=[SaleItemIn(service_id=fijo, precio_cobrado=D(20000))]))
    assert e.value.code == "PRICE_NOT_EDITABLE"
    s = await _sell(p, SaleIn(items=[SaleItemIn(service_id=libre, precio_cobrado=D(35000))], pagos=[EF(35000)]))
    it = s.items[0]
    assert it.precio_original == D("40000") and it.precio_cobrado == D("35000") and it.descuento == D("5000")


async def test_iva_total_and_payment_validation():
    p = await _setup("Básico")
    svc = await _svc(p, "Con IVA", 100000, iva=True)
    await _open_cash(p)
    with pytest.raises(AppError) as e:
        await _sell(p, SaleIn(items=[SaleItemIn(service_id=svc)], propina=D(1000), pagos=[EF(100000)]))
    assert e.value.code == "PAYMENT_MISMATCH" and e.value.details["total"] == "120000.00"
    s = await _sell(p, SaleIn(items=[SaleItemIn(service_id=svc)], propina=D(1000), pagos=[EF(120000)]))
    assert (s.subtotal, s.iva, s.propina, s.total) == (D("100000"), D("19000.00"), D("1000"), D("120000.00"))


async def test_concurrent_collection_of_same_sale_charges_once_and_numbers_have_no_gaps():
    p = await _setup("Básico")
    svc = await _svc(p, "Corte", 10000)
    await _open_cash(p)
    tab = await _sell(p, SaleIn(items=[SaleItemIn(service_id=svc)]))

    async def pay():
        async with scoped_session(tenant_id=p.tenant_id) as db:
            return await sales.pay_sale(db, p, tab.id, PayIn(pagos=[EF(10000)]))
    res = await asyncio.gather(*[pay() for _ in range(5)], return_exceptions=True)
    assert sum(not isinstance(r, Exception) for r in res) == 1, [repr(r)[:200] for r in res if isinstance(r, Exception)]
    assert all(isinstance(r, AppError) and r.code == "CONFLICT" for r in res if isinstance(r, Exception))
    async with scoped_session(tenant_id=p.tenant_id) as db:
        assert (await db.execute(text("SELECT count(*) FROM payments WHERE sale_id = :s"), {"s": tab.id})).scalar() == 1

    async def new_sale(_):
        return await _sell(p, SaleIn(items=[SaleItemIn(service_id=svc)], pagos=[EF(10000)]))
    made = await asyncio.gather(*[new_sale(i) for i in range(6)])
    assert sorted(s.comprobante_numero for s in made) == [f"CI-{n:06d}" for n in range(2, 8)]   # 1 fue la del cobro de arriba


async def test_void_inside_open_cash_releases_appointment_and_excludes_from_cuadre():
    p = await _setup("Básico")
    svc = await _svc(p, "Corte", 30000)
    cid = await _customer(p)
    ses = await _open_cash(p)
    appt = await _appt(p, cid, svc)
    s = await _sell(p, SaleIn(appointment_id=appt.id, pagos=[EF(30000)]))
    async with scoped_session(tenant_id=p.tenant_id) as db:
        v = await sales.void_sale(db, p, s.id, VoidIn(motivo="cliente se arrepintió"))
        assert v.estado == "anulada" and v.anulada_motivo
        assert (await appts.get_appointment(db, p, appt.id)).estado == "en_atencion"
        assert (await cash.cuadre(db, p, ses.id)).ventas_efectivo == D("0")
        with pytest.raises(AppError):
            await sales.void_sale(db, p, s.id, VoidIn(motivo="otra vez"))               # no se anula dos veces
    again = await _sell(p, SaleIn(appointment_id=appt.id, pagos=[EF(30000)]))           # la cita vuelve a poder cobrarse
    assert again.estado == "pagada"


async def test_receipt_and_send_queue():
    p = await _setup("Básico")
    svc = await _svc(p, "Corte", 30000)
    async with scoped_session(tenant_id=p.tenant_id) as db:
        from app.modules.customers import service as cs
        from app.modules.customers.schemas import CustomerIn
        cid = (await cs.create_customer(db, p, CustomerIn(nombre="Ana Ruiz", telefono="3001112233", consentimiento_datos=True))).id
    await _open_cash(p)
    s = await _sell(p, SaleIn(customer_id=cid, items=[SaleItemIn(service_id=svc)], pagos=[EF(30000)]))
    async with scoped_session(tenant_id=p.tenant_id) as db:
        r = await sales.receipt(db, p, s.id)
        out = await sales.send_receipt(db, p, s.id, SendIn(canal="whatsapp"))
        n = (await db.execute(text("SELECT destinatario, estado, plantilla FROM notifications WHERE id = :i"), {"i": out.notificacion_id})).one()
        with pytest.raises(AppError):
            await sales.send_receipt(db, p, s.id, SendIn(canal="email"))                # el cliente no tiene correo
    assert r.numero == s.comprobante_numero and r.cliente == "Ana Ruiz" and r.total == D("30000") and r.moneda == "COP"
    assert (n.destinatario, n.estado, n.plantilla) == ("3001112233", "pendiente", "comprobante")


async def test_reports_match_the_ledger_and_export():
    p = await _setup("Profesional")                                                     # incluye módulo reports
    svc = await _svc(p, "Corte", 30000)
    a, b = await _customer(p), await _customer(p)
    await _open_cash(p)
    s1 = await _sell(p, SaleIn(customer_id=a, items=[SaleItemIn(service_id=svc)], pagos=[EF(30000)]))
    await _sell(p, SaleIn(customer_id=a, items=[SaleItemIn(service_id=svc)], pagos=[PaymentIn(metodo="datafono", monto=D(30000))]))
    await _sell(p, SaleIn(customer_id=b, items=[SaleItemIn(service_id=svc)], propina=D(1000), pagos=[EF(31000)]))
    async with scoped_session(platform=True) as db:                                      # la 1.ª venta de A fue hace 40 días
        await db.execute(text("UPDATE sales SET pagada_en = pagada_en - interval '40 days' WHERE id = :i"), {"i": s1.id})
    hoy = _today()

    async def run(name, d0, d1, agr="dia"):
        async with scoped_session(tenant_id=p.tenant_id) as db:
            return await reports.run(db, p, name, d0, d1, agr, None)

    v = await run("ventas", hoy - timedelta(days=1), hoy + timedelta(days=1))
    assert v.totales["ventas"] == 2 and v.totales["total"] == D("61000") and v.totales["propina"] == D("1000")
    assert v.filas[0]["periodo"] == hoy.isoformat() and v.totales["ticket_promedio"] == D("30500.00")
    todo = await run("ventas", hoy - timedelta(days=60), hoy + timedelta(days=1), "mes")
    assert todo.totales["ventas"] == 3 and todo.totales["total"] == D("91000")

    m = await run("por-metodo", hoy - timedelta(days=60), hoy + timedelta(days=1))
    assert {f["metodo"]: f["monto"] for f in m.filas} == {"efectivo": D("61000"), "datafono": D("30000")}
    assert m.totales["monto"] == todo.totales["total"]                                   # reportes cuadran entre sí
    sv = await run("por-servicio", hoy - timedelta(days=60), hoy + timedelta(days=1))
    assert sv.filas[0]["servicio"] == "Corte" and sv.filas[0]["cantidad"] == 3 and sv.filas[0]["ingresos"] == D("90000")

    cl = await run("clientes", hoy - timedelta(days=1), hoy + timedelta(days=1))
    assert (cl.totales["nuevos"], cl.totales["recurrentes"]) == (1, 1)                  # B nueva; A ya había venido hace 40 días
    cl2 = await run("clientes", hoy - timedelta(days=60), hoy + timedelta(days=1))
    assert (cl2.totales["nuevos"], cl2.totales["recurrentes"]) == (2, 0)

    x = export.to_xlsx(v, "Neg")
    assert load_workbook(BytesIO(x)).active.max_row >= 7 and export.to_pdf(v, "Neg").startswith(b"%PDF")
    with pytest.raises(AppError):
        await run("ventas", hoy - timedelta(days=400), hoy)                              # rango > 366 días
    with pytest.raises(AppError):
        await run("ventas", hoy, hoy - timedelta(days=1))


async def test_appointment_and_occupancy_reports():
    p = await _setup("Profesional")
    st, _ = await _staff_with_schedule(p, "Ana")                                          # 9-18 todos los días = 9 h
    svc = await _svc(p, "Corte", 30000)
    cid = await _customer(p)
    t0 = _tomorrow_10()
    día = t0.astimezone(TZ).date()
    async with scoped_session(tenant_id=p.tenant_id) as db:
        a1 = await appts.create_appointment(db, p, AppointmentIn(customer_id=cid, staff_id=st.id, inicio=t0, duracion_min=60))
        a2 = await appts.create_appointment(db, p, AppointmentIn(customer_id=cid, inicio=t0 + timedelta(hours=3), service_ids=[svc]))
        a3 = await appts.create_appointment(db, p, AppointmentIn(customer_id=cid, inicio=datetime.now(UTC) - timedelta(days=1),
                                                                 service_ids=[svc]))
        await appts.change_status(db, p, a2.id, StatusIn(estado="cancelada", motivo="no puede"))
        await appts.change_status(db, p, a3.id, StatusIn(estado="no_asistio"))
        await appts.change_status(db, p, a1.id, StatusIn(estado="en_atencion"))
        await appts.change_status(db, p, a1.id, StatusIn(estado="atendida"))
        stats = await appts.stats(db, p, datetime.now(UTC) - timedelta(days=3), t0 + timedelta(days=1))   # rango fijo respecto a AHORA: no depende de la hora del día
    assert stats.por_estado == {"atendida": 1, "cancelada": 1, "no_asistio": 1} and stats.total == 3

    async def run(name, d0, d1, agr="dia", staff=None):
        async with scoped_session(tenant_id=p.tenant_id) as db:
            return await reports.run(db, p, name, d0, d1, agr, staff)
    c = await run("citas", datetime.now(TZ).date() - timedelta(days=3), día + timedelta(days=1))   # anclado a hoy: la cita "de ayer" siempre entra
    assert (c.totales["total"], c.totales["atendidas"], c.totales["canceladas"], c.totales["no_asistio"]) == (3, 1, 1, 1)
    assert c.totales["pct_cancelacion"] == 33.3 and c.totales["pct_no_show"] == 33.3

    o = await run("ocupacion", día, día, "dia", st.id)
    assert (o.filas[0]["horas_disponibles"], o.filas[0]["horas_reservadas"], o.filas[0]["ocupacion"]) == (9.0, 1.0, 11.1)
    om = await run("ocupacion", día, día, "mes", st.id)
    assert om.filas[0]["periodo"] == f"{día.year}-{día.month:02d}"
