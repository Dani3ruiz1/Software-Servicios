"""Fase 5 sin BD: dinero, IVA, precio editable, pagos divididos, cuadre, ocupación, exportación y permisos."""
from datetime import UTC, date, datetime, time
from decimal import Decimal
from io import BytesIO
from uuid import uuid4
from zoneinfo import ZoneInfo

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from app.core.errors import AppError
from app.core.security import create_access_token
from app.main import app
from app.modules.cash.service import expected_cash
from app.modules.reports import export
from app.modules.reports import occupancy as occ
from app.modules.reports.schemas import Col, ReportOut
from app.modules.sales.schemas import PaymentIn, SaleIn, SaleItemIn
from app.modules.sales.service import (
    Line,
    item_iva,
    money,
    resolve_price,
    totals,
    validate_payments,
)

client = TestClient(app, raise_server_exceptions=False)
D = Decimal
TZ = ZoneInfo("America/Bogota")


# ---------------------------------------------------------------- dinero
def test_money_rounds_half_up_to_cents():
    assert money(D("10.005")) == D("10.01") and money(D("10.004")) == D("10.00") and money(D("2.5")) == D("2.50")


def test_iva_added_on_top_and_zero_when_not_applicable():
    assert item_iva(D("25000"), True, D("19")) == D("4750.00")
    assert item_iva(D("25000"), False, D("19")) == D("0.00")
    assert item_iva(D("33333"), True, D("19")) == D("6333.27")


def test_totals_match_db_check_total_equals_sum():
    lines = [Line(None, "a", D("100"), D("100"), D("19.00")), Line(None, "b", D("50"), D("40"), D("0.00"))]
    sub, iva, total = totals(lines, D("10"))
    assert (sub, iva, total) == (D("140"), D("19.00"), D("169.00")) and total == sub + iva + D("10")


def test_price_editable_rule_keeps_original():
    assert resolve_price(D("30000"), None, False) == D("30000")
    assert resolve_price(D("30000"), D("30000"), False) == D("30000")
    assert resolve_price(D("30000"), D("25000"), True) == D("25000")
    with pytest.raises(AppError) as e:
        resolve_price(D("30000"), D("25000"), False)
    assert e.value.code == "PRICE_NOT_EDITABLE" and e.value.details["precio_original"] == "30000"


# ---------------------------------------------------------------- pagos
def test_split_payment_must_equal_total():
    ok = [PaymentIn(metodo="efectivo", monto=D("20000")),
          PaymentIn(metodo="transferencia", monto=D("12000"), entidad_bancaria="Bancolombia")]
    validate_payments(D("32000"), ok)
    with pytest.raises(AppError) as e:
        validate_payments(D("32500"), ok)
    assert e.value.code == "PAYMENT_MISMATCH" and e.value.details["diferencia"] == "500"
    with pytest.raises(AppError):
        validate_payments(D("31000"), ok)                                  # de más también falla


def test_transfer_requires_bank_at_schema_and_service_level():
    with pytest.raises(ValidationError):
        PaymentIn(metodo="transferencia", monto=D("1000"))
    with pytest.raises(ValidationError):
        PaymentIn(metodo="transferencia", monto=D("1000"), entidad_bancaria="   ")
    with pytest.raises(ValidationError):
        PaymentIn(metodo="efectivo", monto=D("0"))
    p = PaymentIn.model_construct(metodo="transferencia", monto=D("1000"), entidad_bancaria=None, referencia=None)
    with pytest.raises(AppError) as e:                                      # segunda barrera, por si se salta el schema
        validate_payments(D("1000"), [p])
    assert e.value.code == "BANK_ENTITY_REQUIRED"


def test_sale_input_needs_items_or_appointment_and_valid_items():
    with pytest.raises(ValidationError):
        SaleIn()
    with pytest.raises(ValidationError):
        SaleItemIn()
    with pytest.raises(ValidationError):
        SaleItemIn(descripcion="Gaseosa")                                    # ítem libre exige precio
    assert SaleIn(appointment_id=uuid4()).items == []
    assert SaleItemIn(descripcion="Gaseosa", precio_cobrado=D("3000")).service_id is None


# ---------------------------------------------------------------- caja
def test_cash_expected_ignores_cards_and_transfers():
    assert expected_cash(D("50000"), D("25000"), D("0"), D("3000")) == D("72000")
    assert expected_cash(D("0"), D("0"), D("10000"), D("0")) == D("10000")


# ---------------------------------------------------------------- ocupación
def at(h: int, m: int = 0, day: date = date(2026, 10, 20)) -> datetime:
    return datetime.combine(day, time(h, m), tzinfo=TZ).astimezone(UTC)


def test_occupancy_basic_and_clipped_to_schedule():
    tr = [(time(9), time(13)), (time(14), time(18))]                          # 8 h
    disp, res = occ.staff_day_minutes(date(2026, 10, 20), tr, [], [(at(10), at(11)), (at(12, 30), at(14, 30))], TZ)
    assert disp == 480 and res == 60 + 30 + 30                                # 12:30-13:00 y 14:00-14:30; el almuerzo no cuenta
    assert occ.pct(res, disp) == 25.0


def test_occupancy_blocks_reduce_availability_and_overlapping_blocks_not_double_counted():
    tr = [(time(9), time(17))]
    blocks = [(at(12), at(14)), (at(13), at(15))]                             # se solapan: 12-15 = 3 h
    disp, res = occ.staff_day_minutes(date(2026, 10, 20), tr, blocks, [(at(9), at(10))], TZ)
    assert disp == 480 - 180 and res == 60
    assert occ.pct(0, 0) == 0.0


def test_occupancy_appointment_in_blocked_time_does_not_exceed_100():
    disp, res = occ.staff_day_minutes(date(2026, 10, 20), [(time(9), time(10))], [(at(9), at(10))], [(at(9), at(10))], TZ)
    assert disp == 0 and res == 0


# ---------------------------------------------------------------- exportación
def _report() -> ReportOut:
    return ReportOut(reporte="ventas", titulo="Ventas por dia", desde=date(2026, 10, 1), hasta=date(2026, 10, 31), zona_horaria="America/Bogota",
                     columnas=[Col(key="periodo", label="Periodo", tipo="texto"), Col(key="total", label="Total", tipo="dinero"),
                               Col(key="pct", label="%", tipo="porcentaje")],
                     filas=[{"periodo": "2026-10-01", "total": D("1234567.00"), "pct": 40.5}, {"periodo": "2026-10-02", "total": D("100"), "pct": 59.5}],
                     totales={"periodo": "Total", "total": D("1234667.00"), "pct": 100.0})


def test_cop_format_uses_dots():
    assert export.fmt_cop(D("1234567.00")) == "$ 1.234.567" and export.fmt_cop(0) == "$ 0"


def test_xlsx_export_has_header_rows_and_total():
    wb = load_workbook(BytesIO(export.to_xlsx(_report(), "Barbería Don Pepe")))
    ws = wb.active
    assert ws["A1"].value == "Barbería Don Pepe"
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=5)]
    assert rows[0] == ["Periodo", "Total", "%"] and rows[1][1] == 1234567.0 and rows[-1][0] == "Total" and rows[-1][1] == 1234667.0


def test_pdf_export_is_a_pdf():
    pdf = export.to_pdf(_report(), "Barbería Don Pepe")
    assert pdf.startswith(b"%PDF") and len(pdf) > 1500


# ---------------------------------------------------------------- permisos (sin tocar la BD)
def _h(rol: str) -> dict[str, str]:
    t, _ = create_access_token(sub=uuid4(), kind="user", tenant_id=uuid4(), rol=rol)
    return {"Authorization": f"Bearer {t}"}


def test_only_admin_reads_reports():
    from app.core.rbac import PERMISSIONS as P
    assert "reports:read" in P["admin"] and "reports:read" not in P["recepcionista"] | P["profesional"]


def test_sales_and_cash_permission_matrix():
    from app.core.rbac import PERMISSIONS as P
    assert "sales:void" in P["admin"] and "sales:void" not in P["recepcionista"]
    assert {"sales:read", "sales:write", "cash:manage"} <= P["recepcionista"]
    assert not ({"sales:read", "sales:write", "cash:manage"} & P["profesional"])


def test_new_routes_exist_in_openapi():
    paths = app.openapi()["paths"]
    for p in ("/api/v1/sales", "/api/v1/sales/{sale_id}/pay", "/api/v1/sales/{sale_id}/void", "/api/v1/sales/{sale_id}/receipt",
              "/api/v1/sales/{sale_id}/send", "/api/v1/cash/sessions/open", "/api/v1/cash/sessions/{session_id}/close",
              "/api/v1/cash/sessions/{session_id}/cuadre", "/api/v1/cash/movements", "/api/v1/reports/{nombre}",
              "/api/v1/appointments/stats"):
        assert p in paths, p


def test_report_name_is_a_closed_enum_in_the_contract():
    params = app.openapi()["paths"]["/api/v1/reports/{nombre}"]["get"]["parameters"]
    nombre = next(x for x in params if x["name"] == "nombre")["schema"]
    assert set(nombre["enum"]) == {"ventas", "por-profesional", "por-servicio", "por-metodo", "ocupacion", "clientes", "citas",
                                "ranking-clientes", "clientes-reservas", "reservas-profesional", "reservas-servicio"}
