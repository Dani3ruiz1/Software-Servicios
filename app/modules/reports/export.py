"""Exportación de cualquier ReportOut a Excel (.xlsx) o PDF."""
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.modules.reports.schemas import ReportOut

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF = "application/pdf"


def fmt_cop(v: Any) -> str:
    return "$ " + f"{int(Decimal(str(v)).quantize(Decimal('1'))):,}".replace(",", ".")


def fmt_cell(v: Any, tipo: str) -> str:
    if v is None:
        return ""
    if tipo == "dinero":
        return fmt_cop(v)
    if tipo == "porcentaje":
        return f"{float(v):.1f} %"
    if isinstance(v, float):
        return f"{v:.2f}"
    return str(v)


def to_xlsx(r: ReportOut, negocio: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = r.reporte[:31]
    ws.append([negocio])
    ws.append([r.titulo])
    ws.append([f"Del {r.desde.isoformat()} al {r.hasta.isoformat()} ({r.zona_horaria})"])
    ws.append([])
    head = [c.label for c in r.columnas]
    ws.append(head)
    hr = ws.max_row
    for c in ws[hr]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F5D50")
        c.alignment = Alignment(horizontal="center")
    ws["A1"].font, ws["A2"].font = Font(bold=True, size=14), Font(bold=True)

    def add(row: dict[str, Any], bold: bool = False) -> None:
        ws.append([float(row[c.key]) if isinstance(row.get(c.key), Decimal) else row.get(c.key) for c in r.columnas])
        for i, c in enumerate(r.columnas, start=1):
            cell = ws.cell(row=ws.max_row, column=i)
            if c.tipo == "dinero":
                cell.number_format = '"$" #,##0'
            elif c.tipo == "porcentaje":
                cell.number_format = '0.0"%"'
            if bold:
                cell.font = Font(bold=True)

    for f in r.filas:
        add(f)
    if r.filas:
        add(r.totales, bold=True)
    for i, c in enumerate(r.columnas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(c.label) + 4)
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(r: ReportOut, negocio: str) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=28,
                            title=r.titulo)
    st = getSampleStyleSheet()
    data = [[c.label for c in r.columnas]]
    for f in r.filas:
        data.append([fmt_cell(f.get(c.key), c.tipo) for c in r.columnas])
    if r.filas:
        data.append([fmt_cell(r.totales.get(c.key), c.tipo) for c in r.columnas])
    t = Table(data, repeatRows=1)
    style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F5D50")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
             ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 9),
             ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCCCCC")),
             ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F1EA")]),
             ("ALIGN", (1, 0), (-1, -1), "RIGHT")]
    if r.filas:
        style += [("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("LINEABOVE", (0, -1), (-1, -1), 1, colors.black)]
    t.setStyle(TableStyle(style))
    doc.build([Paragraph(negocio, st["Title"]), Paragraph(r.titulo, st["Heading2"]),
               Paragraph(f"Del {r.desde.isoformat()} al {r.hasta.isoformat()} ({r.zona_horaria})", st["Normal"]),
               Spacer(1, 12), t])
    return buf.getvalue()
